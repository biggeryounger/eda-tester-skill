from __future__ import annotations

import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[3]
CASE_DIR = ROOT / "harness" / "tests" / "e2e" / "optimus_check_design_data_full_workflow"
RUNNER = CASE_DIR / "run_case.py"


def workbook_rows(path: Path) -> list[list[object]]:
    workbook = load_workbook(path, read_only=True, data_only=False)
    rows = [list(row) for row in workbook["cmd"].iter_rows(values_only=True)]
    workbook.close()
    return rows


class FullWorkflowE2ERunnerTests(unittest.TestCase):
    def run_case(self, output: Path) -> subprocess.CompletedProcess[str]:
        return subprocess.run(
            [sys.executable, str(RUNNER), "--output", str(output)],
            cwd=ROOT,
            capture_output=True,
            text=True,
        )

    def test_runner_executes_and_compares_all_expected_outputs(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_dir:
            output = Path(temporary_dir) / "actual"
            completed = self.run_case(output)
            self.assertEqual(0, completed.returncode, completed.stdout + completed.stderr)
            result = json.loads((output / "e2e-result.json").read_text(encoding="utf-8"))
            self.assertEqual("PASS", result["status"])
            self.assertEqual(
                {"structured_spec", "coverage_report", "test_plan", "tcl_cases", "workflow_trace"},
                {check["id"] for check in result["checks"]},
            )
            self.assertTrue(all(check["status"] == "PASS" for check in result["checks"]))
            for relative in (
                "structured-command-spec.json",
                "coverage-report.json",
                "测试用例设计表.xlsx",
                "workflow-trace.json",
                "semantic-review.json",
            ):
                self.assertTrue((output / relative).is_file(), relative)
            self.assertTrue(any((output / "cases").iterdir()))

    def test_runner_is_repeatable_for_normalized_business_outputs(self) -> None:
        with tempfile.TemporaryDirectory() as temporary_dir:
            root = Path(temporary_dir)
            first = root / "first"
            second = root / "second"
            self.assertEqual(0, self.run_case(first).returncode)
            self.assertEqual(0, self.run_case(second).returncode)
            for relative in (
                "structured-command-spec.json",
                "coverage-report.json",
                "workflow-trace.json",
                "semantic-review.json",
                "e2e-result.json",
            ):
                self.assertEqual(
                    (first / relative).read_text(encoding="utf-8"),
                    (second / relative).read_text(encoding="utf-8"),
                    relative,
                )
            self.assertEqual(workbook_rows(first / "测试用例设计表.xlsx"), workbook_rows(second / "测试用例设计表.xlsx"))
            first_tcl = {
                path.relative_to(first / "cases").as_posix(): path.read_text(encoding="utf-8")
                for path in sorted((first / "cases").rglob("*"))
                if path.is_file()
            }
            second_tcl = {
                path.relative_to(second / "cases").as_posix(): path.read_text(encoding="utf-8")
                for path in sorted((second / "cases").rglob("*"))
                if path.is_file()
            }
            self.assertEqual(first_tcl, second_tcl)


if __name__ == "__main__":
    unittest.main()
